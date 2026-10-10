from __future__ import annotations

import argparse
import calendar
import csv
import io
import json
import os
import struct
import sys
import time
import uuid
from contextvars import ContextVar
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable, Iterable

from kakao_pc_lease import PC_KAKAO_UI_LOCK, pc_kakao_ui_lease


DEFAULT_CONFIG_PATH = Path("config/kakao_sender.local.json")

DEFAULT_CONFIG: dict[str, Any] = {
    "search_box": {"x": 0, "y": 0},
    "app_titles": ["KakaoTalk", "카카오톡"],
    "close_chat_key": ["alt", "f4"],
    "clear_search_key": "",
    "clear_search_strategy": "select_text",
    "exit_search_key": "esc",
    "recent_list_top_key": "home",
    "pre_clear_search_key": "",
    "focus_before_each_search": True,
    "restore_main_window_via_uia": True,
    "require_main_window_focus": True,
    "require_window_handles": True,
    "require_new_chat_window": True,
    "verify_chat_title": True,
    "allow_global_close_fallback": False,
    "embedded_chat_close_key": "esc",
    "non_chat_view_close_key": "esc",
    "send_enter_presses": 2,
    "send_enter_interval": 0.2,
    "direct_send_batch_size": 30,
    "direct_send_batch_pause_seconds": 2.0,
    "delays": {
        "after_focus": 0.5,
        "after_click_search": 0.2,
        "after_type_name": 0.5,
        "after_open_chat": 1.0,
        "after_paste_message": 0.7,
        "after_send": 1.0,
        "after_paste_attachments": 1.2,
        "after_send_attachments": 1.4,
        "after_close": 0.8,
        "after_clear_search": 0.4,
        "after_recent_top": 0.2,
        "after_open_schedule_menu": 0.4,
        "after_open_schedule_dialog": 0.8,
        "after_schedule_field_click": 0.2,
        "after_schedule_picker": 0.3,
        "after_schedule_submit": 1.0,
    },
    "press_home_before_search": False,
    "schedule_ui": {
        "plus_button": {"x_offset": 24, "y_offset": -23},
        "plus_button_candidates": [{"x_offset": 36, "y_offset": -42}, {"x_offset": 52, "y_offset": -48}],
        "plus_button_by_chat": {},
        "reservation_menu": {"x_offset": 75, "y_offset": 45},
        "message_field_click": {"x_offset": 176, "y_offset": 170},
        "date_field": {"x_offset": 155, "y_offset": 450},
        "ampm_field": {"x_offset": 270, "y_offset": 450},
        "hour_field": {"x_offset": 288, "y_offset": 450},
        "minute_field": {"x_offset": 321, "y_offset": 450},
        "date_picker": {
            "prev_x_offset": 36,
            "next_x_offset": 206,
            "month_nav_y_offset": 44,
            "grid_x_offset": 29,
            "grid_y_offset": 100,
            "cell_width": 31,
            "cell_height": 28,
        },
        "schedule_button": {"x_offset": 183, "y_offset": 562},
    },
}

NON_CHAT_VIEW_TITLE_MARKERS = ("댓글모아보기",)
SCHEDULE_PLUS_BUTTON_FALLBACK_OFFSETS = ((24.0, -23.0), (36.0, -42.0), (52.0, -48.0))
KAKAO_CHAT_TITLE_SUFFIXES = (" - KakaoTalk", " - 카카오톡")

NAME_ALIASES = {
    "name",
    "chat",
    "chatname",
    "room",
    "roomname",
    "conversation",
    "conversationname",
    "kakaoname",
    "kakaotalkname",
    "카카오톡대화명",
    "카톡대화명",
    "대화명",
    "채팅방",
    "채팅방명",
    "친구",
    "친구명",
}

TEXT_ALIASES = {
    "text",
    "message",
    "body",
    "content",
    "msg",
    "텍스트",
    "내용",
    "메시지",
    "메세지",
    "발송문구",
}

ATTACHMENT_ALIASES = {
    "attachment",
    "attachments",
    "file",
    "files",
    "filepath",
    "filepaths",
    "image",
    "images",
    "photo",
    "photos",
    "첨부",
    "첨부파일",
    "파일",
    "파일경로",
    "이미지",
    "사진",
}

IMAGE_BUNDLE_ALIASES = {
    "imagebundle",
    "bundleimages",
    "bundle",
    "groupimages",
    "이미지묶음",
    "묶어서보내기",
    "묶음",
}

SCHEDULE_ALIASES = {
    "schedule",
    "scheduled",
    "scheduledat",
    "scheduledtime",
    "reservedat",
    "reservation",
    "reservationtime",
    "예약",
    "예약일",
    "예약일시",
    "예약시간",
    "예약메시지",
    "예약메세지",
}

IMAGE_SUFFIXES = {
    ".jpg",
    ".jpeg",
    ".png",
    ".gif",
    ".bmp",
    ".webp",
    ".tif",
    ".tiff",
    ".heic",
    ".heif",
}

ENABLED_ALIASES = {
    "enabled",
    "send",
    "active",
    "use",
    "발송",
    "사용",
    "보내기",
}

FALSE_VALUES = {
    "",
    "0",
    "false",
    "f",
    "n",
    "no",
    "off",
    "skip",
    "disabled",
    "아니오",
    "아니요",
    "미발송",
    "사용안함",
    "건너뜀",
}


@dataclass(frozen=True)
class MessageRow:
    row_number: int
    chat_name: str
    text: str
    attachments: tuple[Path, ...] = ()
    image_bundle: bool = True
    scheduled_at: datetime | None = None

    @property
    def attachment_count(self) -> int:
        return len(self.attachments)

    @property
    def image_count(self) -> int:
        return sum(1 for path in self.attachments if is_image_path(path))


@dataclass(frozen=True)
class OpenChatContext:
    main_handle: int | None
    chat_handle: int | None
    chat_title: str
    process_id: int | None = None


@dataclass(frozen=True, slots=True)
class ScheduleRegistrationResult:
    """Outcome of the registration attempt, separate from verified reservation."""

    verified: bool = False
    detail: str = "예약 등록 버튼 입력 후 창이 닫혔지만 실제 예약 등록 여부는 확인하지 못했습니다."
    acknowledged: bool = False
    cleanup_error: str = ""


class FocusSafetyError(RuntimeError):
    """Raised when the next row could be pasted into the wrong KakaoTalk window."""


class MacroSendCancelled(RuntimeError):
    """The screen operation stopped before a message was dispatched."""


class ScheduleNotSubmittedError(FocusSafetyError):
    """The reservation failed before any registration input was dispatched."""

    def __init__(self, stage: str, cause: BaseException) -> None:
        self.stage = stage
        self.cause = cause
        super().__init__(f"예약 메시지를 등록하기 전에 중단했습니다. ({stage}: {cause})")


class MacroSendUnknownError(FocusSafetyError):
    """A message may have been sent. It must never be retried automatically."""

    def __init__(self, stage: str, cause: BaseException) -> None:
        self.stage = stage
        self.cause = cause
        self.safety_error = isinstance(cause, FocusSafetyError)
        super().__init__(f"전송 결과 확인이 필요합니다. 자동 재전송하지 마세요. ({stage}: {cause})")


@dataclass
class _MacroDispatchState:
    dispatched: bool = False
    stage: str = "prepare"


_MACRO_DISPATCH: ContextVar[_MacroDispatchState | None] = ContextVar("kakao_macro_dispatch", default=None)


def check_macro_cancelled(args: argparse.Namespace) -> None:
    event = getattr(args, "cancel_event", None)
    if event is not None and event.is_set():
        raise MacroSendCancelled("PC 카카오톡 작업이 취소되었습니다.")


def _mark_macro_dispatch(stage: str) -> None:
    state = _MACRO_DISPATCH.get()
    if state is not None:
        state.stage = stage
        state.dispatched = True


def _cancelable_pause(seconds: float, args: argparse.Namespace) -> None:
    event = getattr(args, "cancel_event", None)
    if event is not None:
        if event.wait(seconds):
            check_macro_cancelled(args)
    else:
        time.sleep(seconds)


def normalize_header(value: Any) -> str:
    return (
        str(value or "")
        .strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", "")
    )


def find_column(headers: Iterable[Any], aliases: set[str]) -> str | None:
    for header in headers:
        if normalize_header(header) in aliases:
            return str(header)
    return None


def is_enabled(value: Any) -> bool:
    if value is None:
        return True
    return str(value).strip().lower().replace(" ", "") not in FALSE_VALUES


def parse_bool(value: Any, default: bool = True) -> bool:
    if value is None or str(value).strip() == "":
        return default
    return is_enabled(value)


def parse_schedule_datetime(value: Any) -> datetime | None:
    import re

    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    if isinstance(value, (int, float)) and value > 0:
        try:
            from openpyxl.utils.datetime import from_excel

            parsed = from_excel(value)
            if isinstance(parsed, datetime):
                return parsed
        except Exception:
            pass

    raw = str(value).strip()
    if not raw:
        return None
    raw = raw.replace("T", " ")
    raw = " ".join(raw.split())
    raw = raw.replace("년", "-").replace("월", "-").replace("일", " ")
    raw = raw.replace(".", "-").replace("/", "-")
    raw = raw.replace("시", ":").replace("분", "")
    raw = raw.replace("오전", "AM").replace("오후", "PM")
    raw = re.sub(r"\s*-\s*", "-", raw)
    raw = re.sub(r"\s*:\s*", ":", raw)
    raw = re.sub(r"^(\d{4}-\d{1,2}-\d{1,2})-(?=\d{1,2}:)", r"\1 ", raw)
    raw = re.sub(r"^(\d{4}-\d{1,2}-\d{1,2})-\s+", r"\1 ", raw)
    raw = re.sub(r"^(\d{4}-\d{1,2}-\d{1,2})-$", r"\1", raw)
    raw = " ".join(raw.split())

    korean_ampm_match = None
    if "AM" in raw or "PM" in raw:
        korean_ampm_match = raw

    for fmt in (
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %p %I:%M",
        "%Y-%m-%d %p %I :%M",
        "%Y-%m-%d %p %I: %M",
        "%Y-%m-%d %p %I : %M",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue

    if korean_ampm_match:
        match = re_match_schedule_datetime(korean_ampm_match)
        if match is not None:
            return match

    raise ValueError(f"예약일시 형식을 읽을 수 없습니다: {value!r}. 예: 2026-05-27 09:30")


def re_match_schedule_datetime(value: str) -> datetime | None:
    import re

    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})\s+(AM|PM)\s+(\d{1,2})\s*:?\s*(\d{1,2})$", value)
    if not match:
        return None
    year, month, day, ampm, hour, minute = match.groups()
    hour_value = int(hour)
    if ampm == "PM" and hour_value != 12:
        hour_value += 12
    if ampm == "AM" and hour_value == 12:
        hour_value = 0
    return datetime(int(year), int(month), int(day), hour_value, int(minute))


def format_schedule_datetime(value: datetime | None) -> str:
    if value is None:
        return ""
    return value.strftime("%Y-%m-%d %H:%M")


def split_attachment_value(value: Any) -> list[str]:
    if value is None:
        return []
    raw = str(value).strip()
    if not raw:
        return []
    parts: list[str] = []
    for line in raw.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        for item in line.replace("|", ";").split(";"):
            item = item.strip().strip('"').strip("'")
            if item:
                parts.append(item)
    return parts


def parse_attachment_paths(value: Any, base_dir: Path | None = None) -> tuple[Path, ...]:
    paths: list[Path] = []
    for part in split_attachment_value(value):
        expanded = Path(os.path.expandvars(os.path.expanduser(part)))
        if not expanded.is_absolute() and base_dir is not None:
            expanded = base_dir / expanded
        paths.append(expanded)
    return tuple(paths)


def is_image_path(path: Path) -> bool:
    return path.suffix.lower() in IMAGE_SUFFIXES


def missing_attachment_paths(rows: Iterable[MessageRow]) -> list[tuple[int, Path]]:
    missing: list[tuple[int, Path]] = []
    # One preflight takes one existence observation per path. The transport
    # still validates and snapshots the file before crossing its send boundary.
    exists_by_path: dict[Path, bool] = {}
    for row in rows:
        for path in row.attachments:
            if path not in exists_by_path:
                exists_by_path[path] = path.exists()
            if not exists_by_path[path]:
                missing.append((row.row_number, path))
    return missing


def load_messages(path: Path, sheet: str | None = None) -> list[MessageRow]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        rows = load_csv_messages(path)
    elif suffix == ".xlsx":
        rows = load_xlsx_messages(path, sheet=sheet)
    else:
        raise ValueError("Only .csv and .xlsx files are supported.")

    if not rows:
        raise ValueError("No enabled message rows were found.")
    return rows


def load_csv_messages(path: Path) -> list[MessageRow]:
    data = path.read_bytes()
    if data.startswith(b"PK\x03\x04"):
        return load_xlsx_messages_from_bytes(data, sheet=None, base_dir=path.parent)
    return load_csv_messages_from_bytes(data, source_name=str(path), base_dir=path.parent)


def load_csv_messages_from_bytes(data: bytes, source_name: str = "CSV file", base_dir: Path | None = None) -> list[MessageRow]:
    if not data:
        raise ValueError("CSV file is empty.")

    decode_errors: list[str] = []
    read_errors: list[str] = []
    for encoding in csv_encoding_candidates(data):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError as exc:
            decode_errors.append(f"{encoding}: {exc}")
            continue

        text = text.lstrip("\ufeff")
        if "\x00" in text:
            read_errors.append(f"{encoding}: decoded text still contains NUL bytes")
            continue

        try:
            return load_csv_messages_from_text(text, base_dir=base_dir)
        except (csv.Error, ValueError) as exc:
            read_errors.append(f"{encoding}: {exc}")

    details = "; ".join(read_errors[:3] or decode_errors[:3])
    raise ValueError(
        "Could not read CSV file. Save it as CSV UTF-8, or use the .xlsx file directly. "
        f"Details: {details}"
    )


def csv_encoding_candidates(data: bytes) -> list[str]:
    encodings: list[str] = []
    if data.startswith((b"\xff\xfe", b"\xfe\xff")) or b"\x00" in data[:200]:
        encodings.extend(["utf-16", "utf-16-le", "utf-16-be"])
    encodings.extend(["utf-8-sig", "utf-8", "cp949", "euc-kr"])
    return list(dict.fromkeys(encodings))


def load_csv_messages_from_text(text: str, base_dir: Path | None = None) -> list[MessageRow]:
    stream = io.StringIO(text)
    dialect = sniff_csv_dialect(text)
    reader = csv.reader(stream, dialect=dialect)
    raw_rows = list(reader)
    if not raw_rows:
        raise ValueError("CSV file has no header row.")
    return rows_from_sequence_rows(raw_rows, base_dir=base_dir)


def sniff_csv_dialect(text: str) -> csv.Dialect:
    sample = text[:8192]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        first_line = next((line for line in text.splitlines() if line.strip()), "")
        delimiter = max((",", "\t", ";"), key=lambda candidate: first_line.count(candidate))
        dialect = csv.excel()
        if first_line.count(delimiter) > 0:
            dialect.delimiter = delimiter
        return dialect


def load_xlsx_messages(path: Path, sheet: str | None = None) -> list[MessageRow]:
    return load_xlsx_messages_from_source(path, sheet=sheet, base_dir=path.parent)


def load_xlsx_messages_from_bytes(data: bytes, sheet: str | None = None, base_dir: Path | None = None) -> list[MessageRow]:
    return load_xlsx_messages_from_source(io.BytesIO(data), sheet=sheet, base_dir=base_dir)


def load_xlsx_messages_from_source(source: Any, sheet: str | None = None, base_dir: Path | None = None) -> list[MessageRow]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Reading .xlsx files requires openpyxl. Run: python -m pip install -r requirements.txt"
        ) from exc

    workbook = load_workbook(source, read_only=True, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.active
    raw_rows = list(worksheet.iter_rows(values_only=True))
    if not raw_rows:
        raise ValueError("XLSX file is empty.")

    return rows_from_sequence_rows(raw_rows, base_dir=base_dir)


def rows_from_sequence_rows(raw_rows: list[Iterable[Any]], first_row_number: int = 1, base_dir: Path | None = None) -> list[MessageRow]:
    raw_row_list = [list(row) for row in raw_rows]
    for row_offset, values in enumerate(raw_row_list):
        headers = [str(cell or "").strip().lstrip("\ufeff") for cell in values]
        if find_column(headers, NAME_ALIASES) and (find_column(headers, TEXT_ALIASES) or find_column(headers, ATTACHMENT_ALIASES)):
            return rows_from_header_and_values(headers, raw_row_list[row_offset + 1 :], first_row_number + row_offset + 1, base_dir=base_dir)

    accepted_names = "카카오톡 대화명, 대화명, 채팅방, name, chat"
    accepted_texts = "텍스트, 메시지, 내용, text, message"
    accepted_attachments = "첨부, 첨부파일, 파일경로, 이미지, attachment, image"
    raise ValueError(
        "Required header row was not found. "
        f"Name column examples: {accepted_names}. Text column examples: {accepted_texts}. "
        f"Attachment column examples: {accepted_attachments}."
    )


def rows_from_header_and_values(headers: list[str], raw_rows: list[list[Any]], first_data_row: int, base_dir: Path | None = None) -> list[MessageRow]:
    dict_rows: list[dict[str, Any]] = []
    for values in raw_rows:
        dict_rows.append({headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))})
    return rows_from_dicts(headers, dict_rows, first_data_row=first_data_row, base_dir=base_dir)


def rows_from_dicts(headers: Iterable[Any], rows: list[dict[str, Any]], first_data_row: int, base_dir: Path | None = None) -> list[MessageRow]:
    header_list = [str(header) for header in headers]
    name_column = find_column(header_list, NAME_ALIASES)
    text_column = find_column(header_list, TEXT_ALIASES)
    attachment_column = find_column(header_list, ATTACHMENT_ALIASES)
    image_bundle_column = find_column(header_list, IMAGE_BUNDLE_ALIASES)
    schedule_column = find_column(header_list, SCHEDULE_ALIASES)
    enabled_column = find_column(header_list, ENABLED_ALIASES)

    if not name_column or (not text_column and not attachment_column):
        accepted_names = "카카오톡 대화명, 대화명, 채팅방, name, chat"
        accepted_texts = "텍스트, 메시지, 내용, text, message"
        accepted_attachments = "첨부, 첨부파일, 파일경로, 이미지, attachment, image"
        raise ValueError(
            f"Required columns were not found. Name column examples: {accepted_names}. "
            f"Text column examples: {accepted_texts}. Attachment column examples: {accepted_attachments}."
        )

    result: list[MessageRow] = []
    for index, row in enumerate(rows, start=first_data_row):
        if enabled_column and not is_enabled(row.get(enabled_column)):
            continue
        chat_name = str(row.get(name_column) or "").strip()
        text = str(row.get(text_column) or "").replace("\r\n", "\n").replace("\r", "\n") if text_column else ""
        attachments = parse_attachment_paths(row.get(attachment_column), base_dir=base_dir) if attachment_column else ()
        image_bundle = parse_bool(row.get(image_bundle_column), default=True) if image_bundle_column else True
        try:
            scheduled_at = parse_schedule_datetime(row.get(schedule_column)) if schedule_column else None
        except ValueError as exc:
            raise ValueError(f"Row {index}: {exc}") from exc
        if not chat_name and not text.strip() and not attachments:
            continue
        if not chat_name:
            raise ValueError(f"Row {index}: chat name is empty.")
        if not text.strip() and not attachments:
            raise ValueError(f"Row {index}: message text or attachment is required.")
        result.append(
            MessageRow(
                row_number=index,
                chat_name=chat_name,
                text=text,
                attachments=attachments,
                image_bundle=image_bundle,
                scheduled_at=scheduled_at,
            )
        )
    return result


def deep_merge(base: dict[str, Any], override: dict[str, Any]) -> dict[str, Any]:
    merged = dict(base)
    for key, value in override.items():
        if isinstance(value, dict) and isinstance(merged.get(key), dict):
            merged[key] = deep_merge(merged[key], value)
        else:
            merged[key] = value
    return merged


def normalize_runtime_config(config: dict[str, Any]) -> dict[str, Any]:
    normalized = deep_merge(DEFAULT_CONFIG, config)
    if normalized.get("close_chat_key") == "esc":
        normalized["close_chat_key"] = ["alt", "f4"]
    if normalized.get("clear_search_key") == "esc":
        normalized["clear_search_key"] = ""
    return normalized


def load_config(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"Config file not found: {path}. Run the calibrate command first.")
    with path.open("r", encoding="utf-8") as handle:
        loaded = json.load(handle)
    config = normalize_runtime_config(loaded)
    validate_config(config, path)
    return config


def validate_config(config: dict[str, Any], path: Path) -> None:
    search_box = config.get("search_box") or {}
    x = int(search_box.get("x", 0))
    y = int(search_box.get("y", 0))
    if x <= 0 or y <= 0:
        raise ValueError(f"Search box coordinates are not set in {path}. Run calibrate again.")
    if not config.get("app_titles"):
        raise ValueError("app_titles must contain at least one window title.")


def validate_send_mode(row_count: int, paste_only: bool, no_close_chat: bool) -> None:
    if row_count <= 1:
        return
    if paste_only:
        raise ValueError("붙여넣기만 모드는 여러 행을 연속 처리할 수 없습니다. 1건 테스트로만 사용해 주세요.")
    if no_close_chat:
        raise ValueError("대화창 닫지 않기 모드는 여러 행을 연속 처리할 수 없습니다. 옵션을 끄고 실행해 주세요.")


def save_config(path: Path, config: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(config, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def require_pyautogui():
    try:
        import pyautogui
    except ImportError as exc:
        raise RuntimeError(
            "pyautogui is required for screen control. Run: python -m pip install -r requirements.txt"
        ) from exc
    pyautogui.FAILSAFE = True
    return pyautogui


def win32_window_control_available() -> bool:
    try:
        import win32con  # noqa: F401
        import win32gui  # noqa: F401

        return True
    except Exception:
        return False


def set_clipboard(text: str) -> None:
    try:
        import pyperclip

        pyperclip.copy(text)
        return
    except Exception:
        pass

    try:
        import tkinter

        root = tkinter.Tk()
        root.withdraw()
        root.clipboard_clear()
        root.clipboard_append(text)
        root.update()
        root.destroy()
    except Exception as exc:
        raise RuntimeError(
            "Could not write to clipboard. Install pyperclip with: python -m pip install -r requirements.txt"
        ) from exc


def set_file_clipboard(paths: Iterable[Path]) -> None:
    file_paths = tuple(str(Path(path).resolve()) for path in paths)
    if not file_paths:
        return
    try:
        import win32clipboard
        import win32con
    except Exception as exc:
        raise RuntimeError("파일 첨부에는 pywin32 클립보드 기능이 필요합니다. 프로그램을 다시 설치해 주세요.") from exc

    file_list = "\0".join(file_paths) + "\0\0"
    # CF_HDROP expects a DROPFILES header followed by a double-null-terminated
    # UTF-16 file list. Passing a Python tuple is rejected by pywin32.
    dropfiles_header = struct.pack("<IiiII", 20, 0, 0, 0, 1)
    clipboard_data = dropfiles_header + file_list.encode("utf-16le")

    try:
        win32clipboard.OpenClipboard()
        try:
            win32clipboard.EmptyClipboard()
            win32clipboard.SetClipboardData(win32con.CF_HDROP, clipboard_data)
        finally:
            win32clipboard.CloseClipboard()
    except Exception as exc:
        raise RuntimeError(f"파일 클립보드 설정에 실패했습니다: {exc}") from exc


def send_attachments(pyautogui: Any, config: dict[str, Any], row: MessageRow, args: argparse.Namespace, chat: OpenChatContext | None = None) -> None:
    if not row.attachments:
        return
    missing = [path for path in row.attachments if not path.exists()]
    if missing:
        preview = ", ".join(str(path) for path in missing[:3])
        raise FileNotFoundError(f"첨부 파일을 찾을 수 없습니다: {preview}")

    if chat is None:
        raise FocusSafetyError("첨부를 보낼 대화창의 검증 정보가 없습니다.")
    check_macro_cancelled(args)
    set_file_clipboard(row.attachments)
    verify_window_target(chat.chat_handle, process_id=chat.process_id, expected_title=row.chat_name)
    check_macro_cancelled(args)
    pyautogui.hotkey("ctrl", "v")
    sleep_delay(config, "after_paste_attachments", args.delay_scale, minimum=1.0)

    enter_presses = max(1, int(config.get("attachment_enter_presses", 1)))
    interval = max(0, float(config.get("attachment_enter_interval", 0.2)))
    for press_index in range(enter_presses):
        # Kakao's attachment confirmation may be a child/owned dialog. It must
        # still belong to this chat's process and owner chain, never another app.
        verify_attachment_target(chat)
        check_macro_cancelled(args)
        _mark_macro_dispatch("attachment_dispatch")
        pyautogui.press("enter")
        if press_index < enter_presses - 1:
            _cancelable_pause(interval, args)
    sleep_delay(config, "after_send_attachments", args.delay_scale, minimum=1.0)


def focus_kakao_window(pyautogui: Any, app_titles: list[str]) -> str | None:
    candidates: list[Any] = []
    seen_handles: set[Any] = set()
    for title in app_titles:
        try:
            windows = pyautogui.getWindowsWithTitle(title)
        except Exception:
            windows = []
        for window in windows:
            window_title = str(getattr(window, "title", "") or "")
            handle = getattr(window, "_hWnd", id(window))
            if handle in seen_handles:
                continue
            if not is_usable_kakao_main_window(window, app_titles):
                continue
            candidates.append(window)
            seen_handles.add(handle)

    candidates.sort(key=lambda window: kakao_window_score(window, app_titles), reverse=True)
    for window in candidates:
        try:
            if getattr(window, "isMinimized", False):
                window.restore()
            window.activate()
            return getattr(window, "title", "")
        except Exception:
            handle = getattr(window, "_hWnd", None)
            if handle and activate_window_handle(handle):
                return getattr(window, "title", "")
            continue
    return None


def focus_kakao_window_via_uia(
    app_titles: list[str],
    *,
    desktop_factory: Any | None = None,
    wait_seconds: float = 2.5,
) -> str | None:
    """Restore and focus the PC KakaoTalk main window without saved coordinates.

    Recent Windows/KakaoTalk builds can keep only a taskbar/UIA representation
    while exposing no main window through ``pygetwindow``.  The fallback uses
    the KakaoTalk executable automation ID, never an icon position.  It then
    clicks a small title-bar point and verifies the exact foreground handle
    before returning.  Missing/ambiguous UIA state is a normal safe failure.
    """

    if desktop_factory is None:
        try:
            from pywinauto import Desktop
        except Exception:
            return None
        desktop_factory = Desktop

    expected_titles = {str(title or "").strip().casefold() for title in app_titles if str(title or "").strip()}
    if not expected_titles:
        return None

    def current_main_windows() -> tuple[Any | None, list[Any]]:
        try:
            desktop = desktop_factory(backend="uia")
            windows = []
            for window in desktop.windows():
                try:
                    title = str(window.window_text() or "").strip()
                    info = window.element_info
                    rect = info.rectangle
                    width = int(rect.right) - int(rect.left)
                    height = int(rect.bottom) - int(rect.top)
                    if (
                        title.casefold() in expected_titles
                        and str(info.control_type or "").casefold() == "window"
                        and bool(info.visible)
                        and bool(info.enabled)
                        and width >= 250
                        and height >= 300
                    ):
                        windows.append(window)
                except Exception:
                    continue
            windows.sort(
                key=lambda window: (
                    int(window.element_info.rectangle.right) - int(window.element_info.rectangle.left)
                )
                * (int(window.element_info.rectangle.bottom) - int(window.element_info.rectangle.top)),
                reverse=True,
            )
            return desktop, windows
        except Exception:
            return None, []

    def activate_exact_window(window: Any) -> str | None:
        try:
            info = window.element_info
            rect = info.rectangle
            width = int(rect.right) - int(rect.left)
            height = int(rect.bottom) - int(rect.top)
            if width < 250 or height < 300:
                return None
            # A single click inside the title bar changes focus only; it does
            # not select a friend, chat row, message input, or window control.
            window.click_input(coords=(min(200, max(20, width - 40)), min(15, max(5, height - 1))))
            time.sleep(0.2)
            expected_handle = int(info.handle or 0)
            active_handle = get_foreground_window_handle()
            active_title = get_window_title(active_handle)
            if expected_handle and active_handle == expected_handle and is_usable_kakao_main_title(active_title, app_titles):
                return active_title
        except Exception:
            return None
        return None

    desktop, main_windows = current_main_windows()
    for window in main_windows:
        title = activate_exact_window(window)
        if title:
            return title

    if desktop is None:
        return None
    try:
        taskbar = desktop.window(class_name="Shell_TrayWnd")
        taskbar_buttons = []
        for control in taskbar.descendants():
            try:
                info = control.element_info
                automation_id = str(info.automation_id or "").casefold()
                if (
                    str(info.control_type or "").casefold() == "button"
                    and automation_id.endswith("\\kakao\\kakaotalk\\kakaotalk.exe")
                ):
                    taskbar_buttons.append(control)
            except Exception:
                continue
        if taskbar_buttons:
            # Multiple monitor taskbars may expose the same application
            # button. They all address one KakaoTalk process; click one only.
            taskbar_buttons[0].click_input()
        else:
            # A fully tray-only KakaoTalk has no taskbar application button.
            # Open Windows 11's notification overflow by semantic UIA name,
            # then double-click the exact KakaoTalk notification icon.  No
            # screen coordinate or icon-order assumption is used.
            hidden_icon_buttons = []
            for control in taskbar.descendants():
                try:
                    info = control.element_info
                    label = str(control.window_text() or "").strip().casefold()
                    if (
                        str(info.control_type or "").casefold() == "button"
                        and str(info.automation_id or "").casefold() == "systemtrayicon"
                        and ("숨겨진 아이콘" in label or "hidden icons" in label)
                    ):
                        hidden_icon_buttons.append(control)
                except Exception:
                    continue
            if not hidden_icon_buttons:
                return None
            hidden_icon_buttons[0].click_input()
            time.sleep(0.3)

            overflow_desktop = desktop_factory(backend="uia")
            kakao_tray_icons = []
            for window in overflow_desktop.windows():
                try:
                    if str(window.element_info.class_name or "") != "TopLevelWindowForOverflowXamlIsland":
                        continue
                    for control in window.descendants():
                        info = control.element_info
                        label = str(control.window_text() or "").strip().casefold()
                        if (
                            str(info.control_type or "").casefold() == "button"
                            and str(info.automation_id or "").casefold() == "notifyitemicon"
                            and label in expected_titles
                        ):
                            kakao_tray_icons.append(control)
                except Exception:
                    continue
            if len(kakao_tray_icons) != 1:
                return None
            kakao_tray_icons[0].double_click_input()
    except Exception:
        return None

    deadline = time.monotonic() + max(0.0, min(float(wait_seconds), 5.0))
    while time.monotonic() <= deadline:
        _desktop, main_windows = current_main_windows()
        for window in main_windows:
            title = activate_exact_window(window)
            if title:
                return title
        time.sleep(0.1)
    return None


def is_usable_kakao_main_window(window: Any, app_titles: list[str]) -> bool:
    title = str(getattr(window, "title", "") or "").strip()
    if not is_usable_kakao_main_title(title, app_titles):
        return False
    width = int(getattr(window, "width", 0) or 0)
    height = int(getattr(window, "height", 0) or 0)
    if width < 250 or height < 300:
        return False
    return True


def is_usable_kakao_main_title(title: str, app_titles: list[str]) -> bool:
    title = str(title or "").strip()
    if not title:
        return False
    lowered = title.lower()
    if "shadow" in lowered or "tooltip" in lowered:
        return False
    return any(title == app_title or title.startswith(app_title) for app_title in app_titles)


def normalize_chat_title_for_match(value: str) -> str:
    return "".join(str(value or "").lower().split())


def normalize_kakao_chat_window_title(value: str) -> str:
    title = str(value or "").strip()
    lowered = title.lower()
    for suffix in KAKAO_CHAT_TITLE_SUFFIXES:
        if lowered.endswith(suffix.lower()):
            title = title[: -len(suffix)].rstrip()
            break
    return normalize_chat_title_for_match(title)


def chat_title_matches(expected_chat_name: str, actual_title: str) -> bool:
    expected = normalize_chat_title_for_match(expected_chat_name)
    actual = normalize_kakao_chat_window_title(actual_title)
    if not expected or not actual:
        return False
    return expected == actual


def is_kakao_non_chat_view_title(title: str) -> bool:
    normalized = normalize_chat_title_for_match(title)
    return any(marker in normalized for marker in NON_CHAT_VIEW_TITLE_MARKERS)


def get_root_window_handle_at_point(x: int, y: int) -> int | None:
    try:
        import win32gui

        handle = win32gui.WindowFromPoint((int(x), int(y)))
        if not handle:
            return None
        root = win32gui.GetAncestor(handle, 2)
        return int(root or handle)
    except Exception:
        return None


def search_box_config_from_position(x: int, y: int, app_titles: list[str]) -> dict[str, Any]:
    search_box: dict[str, Any] = {"x": int(x), "y": int(y)}
    handle = get_root_window_handle_at_point(int(x), int(y))
    rect = get_window_rect(handle)
    title = get_window_title(handle)
    if rect is None or not is_usable_kakao_main_title(title, app_titles):
        return search_box

    left, top, right, bottom = rect
    if left <= x <= right and top <= y <= bottom:
        search_box["window_offset"] = {"x": int(x - left), "y": int(y - top)}
    return search_box


def resolve_search_box_position(search_box: dict[str, Any], main_rect: tuple[int, int, int, int] | None) -> tuple[int, int]:
    offset = search_box.get("window_offset")
    if main_rect is not None and isinstance(offset, dict):
        x_offset = int(float(offset.get("x", 0)))
        y_offset = int(float(offset.get("y", 0)))
        return main_rect[0] + x_offset, main_rect[1] + y_offset
    return int(search_box["x"]), int(search_box["y"])


def kakao_window_score(window: Any, app_titles: list[str]) -> tuple[int, int]:
    title = str(getattr(window, "title", "") or "").strip()
    exact = 1 if any(title == app_title for app_title in app_titles) else 0
    area = int(getattr(window, "width", 0) or 0) * int(getattr(window, "height", 0) or 0)
    return exact, area


def get_foreground_window_handle() -> int | None:
    try:
        import win32gui

        handle = int(win32gui.GetForegroundWindow())
        return handle or None
    except Exception:
        return None


def get_window_process_id(handle: int | None) -> int | None:
    if not handle:
        return None
    try:
        import win32process
        _, process_id = win32process.GetWindowThreadProcessId(handle)
        return int(process_id) or None
    except Exception:
        return None


def is_kakao_process_id(process_id: int | None) -> bool:
    if not process_id or os.name != "nt":
        return False
    try:
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.OpenProcess.argtypes = [wintypes.DWORD, wintypes.BOOL, wintypes.DWORD]
        kernel.OpenProcess.restype = wintypes.HANDLE
        kernel.QueryFullProcessImageNameW.argtypes = [wintypes.HANDLE, wintypes.DWORD, wintypes.LPWSTR, ctypes.POINTER(wintypes.DWORD)]
        kernel.QueryFullProcessImageNameW.restype = wintypes.BOOL
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        kernel.CloseHandle.restype = wintypes.BOOL
        process = kernel.OpenProcess(0x1000, False, process_id)
        if not process:
            return False
        try:
            size = wintypes.DWORD(32768)
            buffer = ctypes.create_unicode_buffer(size.value)
            if not kernel.QueryFullProcessImageNameW(process, 0, buffer, ctypes.byref(size)):
                return False
            return Path(buffer.value).name.casefold() == "kakaotalk.exe"
        finally:
            kernel.CloseHandle(process)
    except Exception:
        return False


def verify_window_target(handle: int | None, *, process_id: int | None = None, expected_title: str | None = None, require_foreground: bool = True) -> int:
    """Fail closed immediately before input; return the verified Kakao PID."""
    if not handle or not window_handle_visible(handle):
        raise FocusSafetyError("대상 카카오톡 창이 사라져 작업을 중단했습니다.")
    current_pid = get_window_process_id(handle)
    if not current_pid or (process_id is not None and current_pid != process_id) or not is_kakao_process_id(current_pid):
        raise FocusSafetyError("대상 창의 카카오톡 프로세스를 확인하지 못해 중단했습니다.")
    if require_foreground and get_foreground_window_handle() != handle:
        raise FocusSafetyError("활성 창이 대상 카카오톡 창과 달라 입력을 중단했습니다.")
    if expected_title is not None and (
        not normalize_kakao_chat_window_title(expected_title)
        or normalize_kakao_chat_window_title(expected_title) != normalize_kakao_chat_window_title(get_window_title(handle))
    ):
        raise FocusSafetyError("대상 카카오톡 창 제목이 변경되어 입력을 중단했습니다.")
    return current_pid


def get_window_owner(handle: int) -> int | None:
    try:
        import win32gui
        return int(win32gui.GetWindow(handle, 4)) or None  # GW_OWNER
    except Exception:
        return None


def verify_attachment_target(chat: OpenChatContext) -> None:
    active = get_foreground_window_handle()
    if active == chat.chat_handle:
        verify_window_target(chat.chat_handle, process_id=chat.process_id, expected_title=chat.chat_title)
        return
    verify_window_target(active, process_id=chat.process_id)
    owner = active
    for _ in range(16):
        owner = get_window_owner(owner) if owner else None
        if owner == chat.chat_handle:
            verify_window_target(chat.chat_handle, process_id=chat.process_id, expected_title=chat.chat_title, require_foreground=False)
            return
        if not owner:
            break
    raise FocusSafetyError("첨부 확인 창이 대상 대화창에 속하지 않아 전송을 중단했습니다.")


def _close_cancelled_chat(chat: OpenChatContext | None) -> None:
    """Only close the verified handle; cancellation never focuses/types globally."""
    if chat is None or not chat.chat_handle or chat.chat_handle == chat.main_handle:
        return
    try:
        verify_window_target(chat.chat_handle, process_id=chat.process_id, expected_title=chat.chat_title, require_foreground=False)
        close_window_handle(chat.chat_handle)
    except Exception:
        pass


def activate_window_handle(handle: int) -> bool:
    try:
        import win32con
        import win32gui

        if not win32gui.IsWindow(handle):
            return False
        if win32gui.IsIconic(handle):
            win32gui.ShowWindow(handle, win32con.SW_RESTORE)
        else:
            win32gui.ShowWindow(handle, win32con.SW_SHOW)
        win32gui.SetForegroundWindow(handle)
        return True
    except Exception:
        return False


def window_handle_exists(handle: int | None) -> bool:
    if not handle:
        return False
    try:
        import win32gui

        return bool(win32gui.IsWindow(handle))
    except Exception:
        return False


def window_handle_visible(handle: int | None) -> bool:
    if not handle:
        return False
    try:
        import win32gui

        return bool(win32gui.IsWindow(handle) and win32gui.IsWindowVisible(handle))
    except Exception:
        return False


def wait_until_window_gone(handle: int | None, timeout: float = 1.5) -> bool:
    if not handle:
        return True
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        if not window_handle_visible(handle):
            return True
        time.sleep(0.1)
    return not window_handle_visible(handle)


def close_window_handle(handle: int) -> bool:
    try:
        import win32con
        import win32gui

        if not win32gui.IsWindow(handle):
            return True
        win32gui.PostMessage(handle, win32con.WM_CLOSE, 0, 0)
        return True
    except Exception:
        return False


def close_windows_with_title(title: str, exclude_handles: set[int] | None = None, *, process_id: int | None = None) -> int:
    # Compatibility helper only. Normal sends close their exact handle. Never
    # enumerate close targets without a previously verified Kakao process.
    if process_id is None or not is_kakao_process_id(process_id):
        return 0
    exclude_handles = exclude_handles or set()
    closed = 0
    try:
        import win32gui
    except Exception:
        return 0

    def callback(handle: int, _: Any) -> None:
        nonlocal closed
        if handle in exclude_handles:
            return
        if not win32gui.IsWindowVisible(handle):
            return
        if get_window_process_id(handle) != process_id:
            return
        window_title = win32gui.GetWindowText(handle).strip()
        if window_title == title:
            if close_window_handle(handle):
                closed += 1

    try:
        win32gui.EnumWindows(callback, None)
    except Exception:
        return closed
    return closed


def close_chat_after_send(pyautogui: Any, config: dict[str, Any], args: argparse.Namespace, main_handle: int | None, chat_handle: int | None, chat_title: str, *, process_id: int | None = None) -> None:
    app_titles = list(config.get("app_titles") or [])
    closed_exact_window = False
    if chat_title and is_usable_kakao_main_title(chat_title, app_titles):
        raise FocusSafetyError("대화창 대신 카카오톡 메인 창이 닫힘 대상으로 잡혀 중단했습니다.")
    if not chat_handle:
        raise FocusSafetyError("닫을 대화창 핸들을 확인하지 못해 중단했습니다.")

    verified_pid = verify_window_target(chat_handle, process_id=process_id, expected_title=chat_title, require_foreground=False)

    if chat_handle and chat_handle != main_handle:
        close_window_handle(chat_handle)
        closed_exact_window = wait_until_window_gone(chat_handle, timeout=1.5)
        if not closed_exact_window and activate_window_handle(chat_handle):
            time.sleep(0.1)
            verify_window_target(chat_handle, process_id=verified_pid, expected_title=chat_title)
            press_optional_key(pyautogui, config.get("close_chat_key", ["alt", "f4"]))
            closed_exact_window = wait_until_window_gone(chat_handle, timeout=1.5)

    if closed_exact_window:
        sleep_delay(config, "after_close", args.delay_scale, minimum=0.8)
    elif chat_handle and main_handle and chat_handle == main_handle:
        verify_window_target(chat_handle, process_id=verified_pid, expected_title=chat_title)
        press_optional_key(pyautogui, config.get("embedded_chat_close_key", "esc"))
        sleep_delay(config, "after_close", args.delay_scale, minimum=0.8)
    else:
        raise FocusSafetyError("닫을 대화창 핸들을 확인하지 못해 중단했습니다. 카카오톡 창을 다시 열고 좌표를 확인해 주세요.")

    if chat_handle and chat_handle != main_handle and window_handle_visible(chat_handle):
        raise FocusSafetyError("대화창을 닫지 못해 중단했습니다. 다음 대화명이 채팅 입력창에 붙지 않도록 카카오톡 창 상태를 확인해 주세요.")


def focus_main_and_click_search(pyautogui: Any, config: dict[str, Any], args: argparse.Namespace) -> None:
    check_macro_cancelled(args)
    app_titles = list(config.get("app_titles") or [])
    title = focus_kakao_window(pyautogui, app_titles)
    if not title and config.get("restore_main_window_via_uia", True):
        title = focus_kakao_window_via_uia(app_titles)
    if not title and config.get("require_main_window_focus", True):
        raise FocusSafetyError("카카오톡 메인 창을 찾지 못했습니다. 카카오톡 메인 창을 열어둔 뒤 다시 실행해 주세요.")
    sleep_delay(config, "after_focus", args.delay_scale, minimum=0.5)
    active_handle = get_foreground_window_handle()
    active_title = get_window_title(active_handle)
    if config.get("require_main_window_focus", True) and not is_usable_kakao_main_title(active_title, app_titles):
        raise FocusSafetyError(f"카카오톡 메인 창 대신 '{active_title}' 창이 앞에 있어 중단했습니다. 열린 대화창을 닫고 다시 실행해 주세요.")
    verify_window_target(active_handle)
    check_macro_cancelled(args)
    search_shortcut = config.get("open_search_shortcut")
    if search_shortcut:
        # KakaoTalk's Ctrl+F search entry is anchored to the focused main
        # window, so it remains valid after the window is moved or resized.
        # Keep this opt-in: ordinary calibrated send flows retain their saved
        # coordinate behavior, while read-only matching probes can avoid a
        # stale coordinate without weakening the foreground-title gate above.
        press_optional_key(pyautogui, search_shortcut)
        sleep_delay(config, "after_click_search", args.delay_scale, minimum=0.2)
        return
    search_box = config["search_box"]
    main_rect = get_window_rect(active_handle) if is_usable_kakao_main_title(active_title, app_titles) else None
    x, y = resolve_search_box_position(search_box, main_rect)
    pyautogui.click(x, y)
    sleep_delay(config, "after_click_search", args.delay_scale, minimum=0.2)
    verify_window_target(active_handle)
    check_macro_cancelled(args)


def get_window_title(handle: int | None) -> str:
    if not handle:
        return ""
    try:
        import win32gui

        return win32gui.GetWindowText(handle).strip()
    except Exception:
        return ""


def get_window_rect(handle: int | None) -> tuple[int, int, int, int] | None:
    if not handle:
        return None
    try:
        import win32gui

        if not win32gui.IsWindow(handle):
            return None
        return tuple(int(value) for value in win32gui.GetWindowRect(handle))
    except Exception:
        return None


def get_screen_scale(pyautogui: Any) -> tuple[float, float]:
    try:
        import win32api

        size = pyautogui.size()
        logical_width = max(1, int(win32api.GetSystemMetrics(0)))
        logical_height = max(1, int(win32api.GetSystemMetrics(1)))
        return float(size.width) / logical_width, float(size.height) / logical_height
    except Exception:
        return 1.0, 1.0


def click_logical(pyautogui: Any, x: float, y: float) -> None:
    scale_x, scale_y = get_screen_scale(pyautogui)
    pyautogui.click(int(round(x * scale_x)), int(round(y * scale_y)))


def click_window_offset(pyautogui: Any, rect: tuple[int, int, int, int], x_offset: float, y_offset: float) -> None:
    left, top, right, bottom = rect
    x = left + x_offset if x_offset >= 0 else right + x_offset
    y = top + y_offset if y_offset >= 0 else bottom + y_offset
    click_logical(pyautogui, x, y)


def plus_button_candidate_items(value: Any) -> list[Any]:
    if isinstance(value, dict):
        return [value]
    if isinstance(value, list):
        return value
    return []


def schedule_plus_button_offsets(schedule_ui: Any, chat_title: str = "") -> list[tuple[float, float]]:
    if not isinstance(schedule_ui, dict):
        schedule_ui = {}

    raw_candidates: list[Any] = []
    by_chat = schedule_ui.get("plus_button_by_chat", {})
    if chat_title and isinstance(by_chat, dict):
        raw_candidates.extend(plus_button_candidate_items(by_chat.get(chat_title, [])))
    raw_candidates.append(schedule_ui.get("plus_button", {}))
    raw_candidates.extend(plus_button_candidate_items(schedule_ui.get("plus_button_candidates", [])))
    raw_candidates.extend(
        {"x_offset": x_offset, "y_offset": y_offset}
        for x_offset, y_offset in SCHEDULE_PLUS_BUTTON_FALLBACK_OFFSETS
    )

    offsets: list[tuple[float, float]] = []
    seen: set[tuple[float, float]] = set()
    for candidate in raw_candidates:
        if not isinstance(candidate, dict):
            continue
        try:
            offset = (
                float(candidate.get("x_offset", 24)),
                float(candidate.get("y_offset", -23)),
            )
        except (TypeError, ValueError):
            continue
        if offset in seen:
            continue
        seen.add(offset)
        offsets.append(offset)
    return offsets or [(24.0, -23.0)]


def find_schedule_plus_button_point(
    chat_handle: int, chat_rect: tuple[int, int, int, int],
) -> tuple[int, int] | None:
    """Use a uniquely named attachment control when Kakao exposes one to UIA."""
    try:
        from pywinauto import Desktop

        controls = Desktop(backend="uia").window(handle=chat_handle).descendants(control_type="Button")
    except Exception:
        return None
    left, top, right, bottom = chat_rect
    width, height = right - left, bottom - top
    candidates: set[tuple[int, int]] = set()
    for control in controls:
        try:
            info = control.element_info
            name = str(info.name or "").strip().replace(" ", "")
            rect = info.rectangle
            x = (int(rect.left) + int(rect.right)) // 2
            y = (int(rect.top) + int(rect.bottom)) // 2
            if (
                name in {"+", "첨부", "첨부하기", "파일첨부"}
                and bool(info.visible) and bool(info.enabled)
                and int(rect.right) > int(rect.left) and int(rect.bottom) > int(rect.top)
                and left < x < left + width * 0.35
                and bottom - height * 0.35 < y < bottom
            ):
                candidates.add((x, y))
        except Exception:
            continue
    return next(iter(candidates)) if len(candidates) == 1 else None


def find_kakao_windows(
    title: str | None = None,
    class_name: str | None = None,
    exclude_handles: set[int] | None = None,
    min_size: tuple[int, int] = (20, 20),
) -> list[tuple[int, str, str, tuple[int, int, int, int]]]:
    exclude_handles = exclude_handles or set()
    results: list[tuple[int, str, str, tuple[int, int, int, int]]] = []
    try:
        import win32gui
        import win32process
    except Exception:
        return results

    def callback(handle: int, _: Any) -> None:
        if handle in exclude_handles:
            return
        if not win32gui.IsWindowVisible(handle):
            return
        try:
            _, process_id = win32process.GetWindowThreadProcessId(handle)
            if not is_kakao_process_id(process_id):
                return
            window_title = win32gui.GetWindowText(handle).strip()
            window_class = win32gui.GetClassName(handle)
            rect = tuple(int(value) for value in win32gui.GetWindowRect(handle))
        except Exception:
            return
        width = rect[2] - rect[0]
        height = rect[3] - rect[1]
        if width < min_size[0] or height < min_size[1]:
            return
        if class_name and window_class != class_name:
            return
        if title is not None and window_title != title:
            return
        if not (window_class.startswith("EVA") or "Kakao" in window_class or window_title in {"카카오톡", "예약 메시지", "브리핑 보드"}):
            return
        results.append((handle, window_class, window_title, rect))

    try:
        win32gui.EnumWindows(callback, None)
    except Exception:
        return results
    return results


def wait_for_kakao_window(
    title: str | None = None,
    class_name: str | None = None,
    exclude_handles: set[int] | None = None,
    timeout: float = 2.0,
) -> tuple[int, str, str, tuple[int, int, int, int]] | None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        matches = find_kakao_windows(title=title, class_name=class_name, exclude_handles=exclude_handles)
        if matches:
            matches.sort(key=lambda item: (item[3][2] - item[3][0]) * (item[3][3] - item[3][1]), reverse=True)
            return matches[0]
        time.sleep(0.05)
    return None


def get_child_rect(parent_handle: int, class_name: str) -> tuple[int, int, int, int] | None:
    try:
        import win32gui
    except Exception:
        return None
    result: tuple[int, int, int, int] | None = None

    def callback(child_handle: int, _: Any) -> None:
        nonlocal result
        if result is not None:
            return
        try:
            if win32gui.GetClassName(child_handle) == class_name and win32gui.IsWindowVisible(child_handle):
                result = tuple(int(value) for value in win32gui.GetWindowRect(child_handle))
        except Exception:
            return

    try:
        win32gui.EnumChildWindows(parent_handle, callback, None)
    except Exception:
        return None
    return result


def countdown(seconds: int, label: str) -> None:
    if seconds <= 0:
        return
    for remaining in range(seconds, 0, -1):
        print(f"{label}: {remaining}...", end="\r", flush=True)
        time.sleep(1)
    print(" " * 40, end="\r")


def sleep_delay(config: dict[str, Any], key: str, scale: float, minimum: float = 0) -> None:
    if scale <= 0:
        return
    seconds = max(float(config.get("delays", {}).get(key, 0)), minimum) * scale
    if seconds > 0:
        time.sleep(seconds)


def press_optional_key(pyautogui: Any, key: Any) -> None:
    if not key:
        return
    if isinstance(key, list):
        pyautogui.hotkey(*key)
        return
    pyautogui.press(str(key))


def clear_search_box(pyautogui: Any, config: dict[str, Any]) -> None:
    strategy = str(config.get("clear_search_strategy") or "select_text")
    if strategy == "none":
        return
    if strategy == "key":
        press_optional_key(pyautogui, config.get("clear_search_key"))
        return
    if strategy == "backspace":
        for _ in range(int(config.get("clear_search_backspaces", 80))):
            pyautogui.press("backspace")
        return
    pyautogui.press("home")
    pyautogui.hotkey("shift", "end")
    pyautogui.press("backspace")


def return_to_recent_chat_top(pyautogui: Any, config: dict[str, Any], args: argparse.Namespace) -> None:
    focus_main_and_click_search(pyautogui, config, args)
    clear_search_box(pyautogui, config)
    press_optional_key(pyautogui, config.get("exit_search_key", "esc"))
    sleep_delay(config, "after_clear_search", args.delay_scale, minimum=0.2)
    pyautogui.press(str(config.get("recent_list_top_key") or "home"))
    sleep_delay(config, "after_recent_top", args.delay_scale, minimum=0.2)


def prepare_kakao(pyautogui: Any, config: dict[str, Any], args: argparse.Namespace) -> None:
    if args.assume_focused:
        return

    title = focus_kakao_window(pyautogui, list(config["app_titles"]))
    if title:
        print(f"Activated KakaoTalk window: {title}")
        sleep_delay(config, "after_focus", args.delay_scale)
        return

    print("Could not activate KakaoTalk automatically.")
    print("Make KakaoTalk visible and keep the calibrated search box position unobstructed.")
    input("Press Enter here, then switch to KakaoTalk during the countdown.")
    countdown(args.countdown, "Starting")


def open_chat(pyautogui: Any, config: dict[str, Any], row: MessageRow, args: argparse.Namespace) -> OpenChatContext:
    check_macro_cancelled(args)
    prepare_kakao(pyautogui, config, args)
    if config.get("require_main_window_focus", True) and not win32_window_control_available():
        raise FocusSafetyError("창 핸들 제어 모듈(pywin32)을 사용할 수 없어 중단했습니다. 프로그램을 다시 설치해 주세요.")

    if config.get("press_home_before_search"):
        verify_window_target(get_foreground_window_handle())
        check_macro_cancelled(args)
        pyautogui.press("home")

    if config.get("focus_before_each_search", True):
        focus_main_and_click_search(pyautogui, config, args)
    else:
        search_box = config["search_box"]
        main_rect = None
        if isinstance(search_box.get("window_offset"), dict):
            active_handle = get_foreground_window_handle()
            active_title = get_window_title(active_handle)
            app_titles = list(config.get("app_titles") or [])
            main_rect = get_window_rect(active_handle) if is_usable_kakao_main_title(active_title, app_titles) else None
        x, y = resolve_search_box_position(search_box, main_rect)
        pyautogui.click(x, y)
        sleep_delay(config, "after_click_search", args.delay_scale, minimum=0.2)

    main_handle = get_foreground_window_handle()
    if config.get("require_window_handles", True) and not main_handle:
        raise FocusSafetyError("카카오톡 메인 창 핸들을 확인하지 못해 중단했습니다. 프로그램을 다시 실행해 주세요.")

    main_pid = verify_window_target(main_handle)
    check_macro_cancelled(args)
    press_optional_key(pyautogui, config.get("pre_clear_search_key"))
    clear_search_box(pyautogui, config)
    set_clipboard(row.chat_name)
    verify_window_target(main_handle, process_id=main_pid)
    check_macro_cancelled(args)
    pyautogui.hotkey("ctrl", "v")
    sleep_delay(config, "after_type_name", args.delay_scale, minimum=0.5)

    verify_window_target(main_handle, process_id=main_pid)
    check_macro_cancelled(args)
    pyautogui.press("enter")
    sleep_delay(config, "after_open_chat", args.delay_scale, minimum=1.0)
    chat_handle = get_foreground_window_handle()
    chat_title = get_window_title(chat_handle)
    app_titles = list(config.get("app_titles") or [])
    if config.get("require_window_handles", True) and not chat_handle:
        raise FocusSafetyError("대화창 핸들을 확인하지 못해 중단했습니다. 카카오톡 대화창이 정상적으로 열렸는지 확인해 주세요.")
    if chat_title and is_kakao_non_chat_view_title(chat_title):
        press_optional_key(pyautogui, config.get("non_chat_view_close_key", "esc"))
        sleep_delay(config, "after_close", args.delay_scale, minimum=0.2)
        raise FocusSafetyError(
            f"채팅방이 아닌 '{chat_title}' 화면이 열려 중단했습니다. "
            "카카오톡 검색창 좌표를 다시 찍은 뒤 재시도해 주세요."
        )
    if chat_title and is_usable_kakao_main_title(chat_title, app_titles):
        raise FocusSafetyError("대화창 대신 카카오톡 메인 창이 활성화되어 중단했습니다. 대화명이 정확한지 확인해 주세요.")
    if (
        config.get("require_new_chat_window", True)
        and main_handle
        and chat_handle
        and chat_handle == main_handle
    ):
        raise FocusSafetyError("대화창이 새 창으로 열렸는지 확인하지 못해 중단했습니다. 대화명이 정확한지와 카카오톡 검색창 좌표를 확인해 주세요.")
    if config.get("verify_chat_title", True):
        if not chat_title:
            raise FocusSafetyError("열린 대화창 제목을 확인하지 못해 중단했습니다. 대화명이 정확한지 확인해 주세요.")
        if not chat_title_matches(row.chat_name, chat_title):
            raise FocusSafetyError(f"열린 대화창 제목('{chat_title}')이 대상 대화명('{row.chat_name}')과 달라 중단했습니다.")

    chat_pid = verify_window_target(chat_handle, process_id=main_pid, expected_title=row.chat_name)
    check_macro_cancelled(args)
    return OpenChatContext(main_handle=main_handle, chat_handle=chat_handle, chat_title=chat_title, process_id=chat_pid)


def send_one(pyautogui: Any, config: dict[str, Any], row: MessageRow, args: argparse.Namespace) -> None:
    with pc_kakao_ui_lease(cancel_event=getattr(args, "cancel_event", None), timeout=getattr(args, "pc_lease_timeout", 30.0)):
        state = _MacroDispatchState()
        token = _MACRO_DISPATCH.set(state)
        chat = None
        try:
            check_macro_cancelled(args)
            missing = missing_attachment_paths([row])
            if missing:
                raise FileNotFoundError("첨부 파일을 찾을 수 없습니다: " + ", ".join(str(path) for _, path in missing[:3]))
            if args.paste_only and not row.text.strip():
                raise ValueError("붙여넣기만 모드는 텍스트가 필요합니다.")
            chat = open_chat(pyautogui, config, row, args)
            if row.text.strip():
                state.stage = "text_paste"
                check_macro_cancelled(args)
                set_clipboard(row.text)
                verify_window_target(chat.chat_handle, process_id=chat.process_id, expected_title=row.chat_name)
                check_macro_cancelled(args)
                pyautogui.hotkey("ctrl", "v")
                sleep_delay(config, "after_paste_message", args.delay_scale, minimum=0.7)
                check_macro_cancelled(args)
                if args.paste_only:
                    return
                presses = max(1, int(config.get("send_enter_presses", 2)))
                interval = max(0, float(config.get("send_enter_interval", 0.2)))
                for press_index in range(presses):
                    verify_window_target(chat.chat_handle, process_id=chat.process_id, expected_title=row.chat_name)
                    check_macro_cancelled(args)
                    _mark_macro_dispatch("text_dispatch")
                    pyautogui.press("enter")
                    if press_index < presses - 1:
                        _cancelable_pause(interval, args)
                sleep_delay(config, "after_send", args.delay_scale, minimum=1.0)
            if row.attachments and not args.paste_only:
                state.stage = "attachment_paste"
                send_attachments(pyautogui, config, row, args, chat)
            state.stage = "cleanup"
            check_macro_cancelled(args)
            if not args.no_close_chat:
                close_chat_after_send(pyautogui, config, args, chat.main_handle, chat.chat_handle, chat.chat_title, process_id=chat.process_id)
                focus_main_and_click_search(pyautogui, config, args)
                clear_search_box(pyautogui, config)
                sleep_delay(config, "after_clear_search", args.delay_scale, minimum=0.4)
        except (Exception, KeyboardInterrupt) as exc:
            if isinstance(exc, (MacroSendCancelled, KeyboardInterrupt)):
                _close_cancelled_chat(chat)
            if state.dispatched:
                raise MacroSendUnknownError(state.stage, exc) from exc
            raise
        finally:
            _MACRO_DISPATCH.reset(token)


def schedule_datetime_for_row(row: MessageRow, fallback: datetime | None = None) -> datetime | None:
    return row.scheduled_at or fallback


def earliest_schedule_datetime(now: datetime) -> datetime:
    """Return the first five-minute slot at least five full minutes from now."""
    threshold = now + timedelta(minutes=5)
    minute_floor = threshold.replace(second=0, microsecond=0)
    candidate = minute_floor + timedelta(minutes=(-minute_floor.minute) % 5)
    if candidate < threshold:
        candidate += timedelta(minutes=5)
    return candidate


def validate_schedule_row(row: MessageRow, scheduled_at: datetime | None, now: datetime | None = None) -> None:
    if scheduled_at is None:
        raise ValueError(f"행 {row.row_number}: 예약일시가 필요합니다.")
    if scheduled_at.minute % 5:
        raise ValueError(f"행 {row.row_number}: 카카오톡 예약 시각의 분은 5분 단위여야 합니다.")
    if scheduled_at.second or scheduled_at.microsecond:
        raise ValueError(f"행 {row.row_number}: 카카오톡 예약 시각의 초와 마이크로초는 0이어야 합니다.")
    if not row.text.strip():
        raise ValueError(f"행 {row.row_number}: 예약 메시지는 텍스트가 필요합니다.")
    if row.attachments:
        raise ValueError(f"행 {row.row_number}: 예약 메시지는 현재 텍스트만 지원합니다. 첨부 파일은 제거해 주세요.")
    now = now or datetime.now()
    if scheduled_at < earliest_schedule_datetime(now):
        raise ValueError(f"행 {row.row_number}: 예약일시는 현재 시각에서 최소 5분 이후여야 합니다.")


def send_scheduled_one(
    pyautogui: Any,
    config: dict[str, Any],
    row: MessageRow,
    args: argparse.Namespace,
    scheduled_at: datetime,
) -> ScheduleRegistrationResult:
    state = _MacroDispatchState(stage="schedule_lease")
    token = _MACRO_DISPATCH.set(state)
    chat = None
    try:
        with pc_kakao_ui_lease(cancel_event=getattr(args, "cancel_event", None), timeout=getattr(args, "pc_lease_timeout", 30.0)):
            try:
                state.stage = "prepare"
                check_macro_cancelled(args)
                validate_schedule_row(row, scheduled_at)
                state.stage = "chat_open"
                chat = open_chat(pyautogui, config, row, args)
                state.stage = "schedule_dialog"
                schedule_handle = open_schedule_dialog(pyautogui, config, args, chat)
                state.stage = "schedule_form"
                schedule_rect = get_window_rect(schedule_handle)
                if schedule_rect is None:
                    raise FocusSafetyError("예약 메시지 창 위치를 확인하지 못했습니다.")
                verify_window_target(schedule_handle, process_id=chat.process_id, expected_title="예약 메시지")
                paste_schedule_message(pyautogui, config, args, schedule_handle, schedule_rect, row.text)
                set_schedule_datetime(pyautogui, config, args, schedule_handle, schedule_rect, scheduled_at)
                # The native form does not yet expose a verified readback port.
                # Clock-derived picker navigation only prepares it; a person must
                # compare the displayed recipient, body and full date/time.
                state.stage = "schedule_review"
                confirm_ready = getattr(args, "confirm_schedule_ready", None)
                if not callable(confirm_ready):
                    raise FocusSafetyError("예약 메시지 입력값을 확인할 사용자 확인 절차가 없어 등록하지 않았습니다.")
                check_macro_cancelled(args)
                if confirm_ready(row, scheduled_at) is not True:
                    raise MacroSendCancelled("사용자가 예약 입력값을 확인하지 않아 등록 전에 중단했습니다.")
                state.stage = "schedule_submit"
                check_macro_cancelled(args)
                if not activate_window_handle(schedule_handle):
                    raise FocusSafetyError("확인한 예약 메시지 창을 앞으로 가져오지 못해 등록하지 않았습니다.")
                verify_window_target(schedule_handle, process_id=chat.process_id, expected_title="예약 메시지")
                # Time spent opening/editing the native form may consume the lead
                # time. Reject before entering the irreversible submit helper.
                validate_schedule_row(row, scheduled_at)
                click_schedule_submit(
                    pyautogui, config, args, schedule_handle, schedule_rect,
                    before_dispatch=lambda: validate_schedule_row(row, scheduled_at),
                )
                # The real submit helper marks before clicking, including uncertain
                # input exceptions. A returned injected helper also means dispatched.
                state.dispatched = True
                if not wait_until_window_gone(schedule_handle, timeout=3.0):
                    raise FocusSafetyError("예약 메시지 창이 닫히지 않아 예약 등록 여부를 확인하지 못했습니다.")
                state.stage = "cleanup"
                check_macro_cancelled(args)
                if not args.no_close_chat:
                    close_chat_after_send(pyautogui, config, args, chat.main_handle, chat.chat_handle, chat.chat_title, process_id=chat.process_id)
                    focus_main_and_click_search(pyautogui, config, args)
                    clear_search_box(pyautogui, config)
                    sleep_delay(config, "after_clear_search", args.delay_scale, minimum=0.4)
                return ScheduleRegistrationResult()
            except (MacroSendCancelled, KeyboardInterrupt):
                # Exact-window cancellation cleanup must still own the UI lease.
                _close_cancelled_chat(chat)
                raise
    except (Exception, KeyboardInterrupt) as exc:
        # Include lease acquisition and release: an exception while releasing
        # the UI after submission cannot turn into a safely retryable failure.
        if state.dispatched:
            raise MacroSendUnknownError(state.stage, exc) from exc
        if isinstance(exc, (MacroSendCancelled, KeyboardInterrupt)):
            raise
        raise ScheduleNotSubmittedError(state.stage, exc) from exc
    finally:
        _MACRO_DISPATCH.reset(token)


def open_schedule_dialog(pyautogui: Any, config: dict[str, Any], args: argparse.Namespace, chat: OpenChatContext) -> int:
    check_macro_cancelled(args)
    if not chat.chat_handle:
        raise FocusSafetyError("예약 메시지를 열 대화창 핸들을 확인하지 못했습니다.")
    if not activate_window_handle(chat.chat_handle):
        raise FocusSafetyError("예약 메시지를 열 대화창을 앞으로 가져오지 못했습니다.")
    verify_window_target(chat.chat_handle, process_id=chat.process_id, expected_title=chat.chat_title)
    chat_rect = get_window_rect(chat.chat_handle)
    if chat_rect is None:
        raise FocusSafetyError("대화창 위치를 확인하지 못했습니다.")

    schedule_ui = config.get("schedule_ui", {})
    if not isinstance(schedule_ui, dict):
        schedule_ui = {}
    existing = {handle for handle, _, _, _ in find_kakao_windows()}

    plus_point = find_schedule_plus_button_point(chat.chat_handle, chat_rect)
    plus_offsets = schedule_plus_button_offsets(schedule_ui, chat.chat_title)
    menu = None
    if plus_point is not None:
        check_macro_cancelled(args)
        verify_window_target(chat.chat_handle, process_id=chat.process_id, expected_title=chat.chat_title)
        click_logical(pyautogui, *plus_point)
        sleep_delay(config, "after_open_schedule_menu", args.delay_scale, minimum=0.2)
        menu = wait_for_kakao_window(class_name="EVA_Menu", exclude_handles=existing, timeout=2.0)
    for x_offset, y_offset in plus_offsets:
        if menu is not None:
            break
        check_macro_cancelled(args)
        verify_window_target(chat.chat_handle, process_id=chat.process_id, expected_title=chat.chat_title)
        click_window_offset(pyautogui, chat_rect, x_offset, y_offset)
        sleep_delay(config, "after_open_schedule_menu", args.delay_scale, minimum=0.2)
        menu = wait_for_kakao_window(class_name="EVA_Menu", exclude_handles=existing, timeout=2.0)
        if menu is not None:
            break
    if menu is None:
        tried = ", ".join(f"({x_offset:g}, {y_offset:g})" for x_offset, y_offset in plus_offsets)
        raise FocusSafetyError(
            f"카카오톡 + 메뉴를 찾지 못했습니다. schedule_ui.plus_button 또는 "
            f"schedule_ui.plus_button_by_chat 좌표를 확인해 주세요. 시도 좌표: {tried}"
        )
    menu_handle, _, _, menu_rect = menu
    verify_window_target(menu_handle, process_id=chat.process_id)
    check_macro_cancelled(args)
    reservation = schedule_ui.get("reservation_menu", {})
    click_window_offset(
        pyautogui,
        menu_rect,
        float(reservation.get("x_offset", 75)),
        float(reservation.get("y_offset", 45)),
    )
    sleep_delay(config, "after_open_schedule_dialog", args.delay_scale, minimum=0.5)

    schedule_window = wait_for_kakao_window(title="예약 메시지", exclude_handles=existing, timeout=3.0)
    if schedule_window is None:
        raise FocusSafetyError("예약 메시지 창을 찾지 못했습니다.")
    verify_window_target(schedule_window[0], process_id=chat.process_id, expected_title="예약 메시지")
    return int(schedule_window[0])


def paste_schedule_message(
    pyautogui: Any,
    config: dict[str, Any],
    args: argparse.Namespace,
    schedule_handle: int,
    schedule_rect: tuple[int, int, int, int],
    text: str,
) -> None:
    check_macro_cancelled(args)
    process_id = verify_window_target(schedule_handle, expected_title="예약 메시지")
    message_rect = get_child_rect(schedule_handle, "RICHEDIT50W")
    if message_rect is not None:
        click_logical(pyautogui, (message_rect[0] + message_rect[2]) / 2, (message_rect[1] + message_rect[3]) / 2)
    else:
        field = config.get("schedule_ui", {}).get("message_field_click", {})
        click_window_offset(
            pyautogui,
            schedule_rect,
            float(field.get("x_offset", 176)),
            float(field.get("y_offset", 170)),
        )
    set_clipboard(text)
    verify_window_target(schedule_handle, process_id=process_id, expected_title="예약 메시지")
    check_macro_cancelled(args)
    pyautogui.hotkey("ctrl", "v")
    sleep_delay(config, "after_paste_message", args.delay_scale, minimum=0.4)


def set_schedule_datetime(
    pyautogui: Any,
    config: dict[str, Any],
    args: argparse.Namespace,
    schedule_handle: int,
    schedule_rect: tuple[int, int, int, int],
    scheduled_at: datetime,
) -> None:
    if scheduled_at.date() != datetime.now().date():
        select_schedule_date(pyautogui, config, args, schedule_handle, schedule_rect, scheduled_at.date())
    select_schedule_time(pyautogui, config, args, schedule_rect, scheduled_at)


def select_schedule_date(
    pyautogui: Any,
    config: dict[str, Any],
    args: argparse.Namespace,
    schedule_handle: int,
    schedule_rect: tuple[int, int, int, int],
    target_date: date,
) -> None:
    check_macro_cancelled(args)
    process_id = verify_window_target(schedule_handle, expected_title="예약 메시지")
    schedule_ui = config.get("schedule_ui", {})
    date_field = schedule_ui.get("date_field", {})
    existing = {handle for handle, _, _, _ in find_kakao_windows()}
    click_window_offset(pyautogui, schedule_rect, float(date_field.get("x_offset", 155)), float(date_field.get("y_offset", 450)))
    sleep_delay(config, "after_schedule_field_click", args.delay_scale, minimum=0.2)
    picker = wait_for_kakao_window(class_name="EVA_Window_Dblclk", exclude_handles=existing | {schedule_handle}, timeout=2.0)
    if picker is None:
        raise FocusSafetyError("예약 날짜 선택 창을 찾지 못했습니다.")
    verify_window_target(picker[0], process_id=process_id)
    picker_rect = picker[3]
    date_picker = schedule_ui.get("date_picker", {})
    current_month = datetime.now().replace(day=1).date()
    target_month = target_date.replace(day=1)
    month_delta = (target_month.year - current_month.year) * 12 + (target_month.month - current_month.month)
    nav_x_key = "next_x_offset" if month_delta >= 0 else "prev_x_offset"
    for _ in range(abs(month_delta)):
        check_macro_cancelled(args)
        verify_window_target(picker[0], process_id=process_id)
        click_window_offset(
            pyautogui,
            picker_rect,
            float(date_picker.get(nav_x_key, 206 if month_delta >= 0 else 36)),
            float(date_picker.get("month_nav_y_offset", 44)),
        )
        sleep_delay(config, "after_schedule_picker", args.delay_scale, minimum=0.15)

    calendar_rows = calendar.Calendar(firstweekday=6).monthdayscalendar(target_date.year, target_date.month)
    row_index = next((index for index, row in enumerate(calendar_rows) if target_date.day in row), None)
    if row_index is None:
        raise FocusSafetyError("예약 날짜 위치를 계산하지 못했습니다.")
    col_index = calendar_rows[row_index].index(target_date.day)
    x_offset = float(date_picker.get("grid_x_offset", 29)) + col_index * float(date_picker.get("cell_width", 31))
    y_offset = float(date_picker.get("grid_y_offset", 100)) + row_index * float(date_picker.get("cell_height", 28))
    verify_window_target(picker[0], process_id=process_id)
    check_macro_cancelled(args)
    click_window_offset(pyautogui, picker_rect, x_offset, y_offset)
    sleep_delay(config, "after_schedule_picker", args.delay_scale, minimum=0.2)


def select_schedule_time(
    pyautogui: Any,
    config: dict[str, Any],
    args: argparse.Namespace,
    schedule_rect: tuple[int, int, int, int],
    scheduled_at: datetime,
) -> None:
    schedule_ui = config.get("schedule_ui", {})
    if not isinstance(schedule_ui, dict):
        schedule_ui = {}
    default_time = default_schedule_time(datetime.now())
    default_is_pm = default_time.hour >= 12
    target_is_pm = scheduled_at.hour >= 12

    if target_is_pm != default_is_pm:
        _, _, _, ampm_rect = open_schedule_time_popup(pyautogui, schedule_ui, config, args, schedule_rect, "ampm_field")
        left, top, right, bottom = ampm_rect
        if bottom - top < 50:
            raise FocusSafetyError("예약 시각의 오전 선택 항목을 찾지 못했습니다. 대상 날짜와 시각을 확인해 주세요.")
        click_logical(pyautogui, (left + right) / 2, top + (bottom - top) * (0.75 if target_is_pm else 0.25))
        sleep_delay(config, "after_schedule_picker", args.delay_scale, minimum=0.2)

    default_hour = default_time.hour % 12 or 12
    target_hour = scheduled_at.hour % 12 or 12
    for field_name, current_index, target_index in (
        ("hour_field", default_hour - 1, target_hour - 1),
        ("minute_field", default_time.minute // 5, scheduled_at.minute // 5),
    ):
        if current_index == target_index:
            continue
        popup = open_schedule_time_popup(pyautogui, schedule_ui, config, args, schedule_rect, field_name)
        popup_pid = verify_window_target(popup[0])
        key, presses = cyclic_dropdown_key_steps(current_index, target_index, 12)
        for _ in range(presses):
            check_macro_cancelled(args)
            verify_window_target(popup[0], process_id=popup_pid)
            pyautogui.press(key)
        verify_window_target(popup[0], process_id=popup_pid)
        check_macro_cancelled(args)
        pyautogui.press("enter")
        sleep_delay(config, "after_schedule_picker", args.delay_scale, minimum=0.2)


def default_schedule_time(now: datetime) -> datetime:
    return earliest_schedule_datetime(now)


def cyclic_dropdown_key_steps(current_index: int, target_index: int, option_count: int) -> tuple[str, int]:
    forward_steps = (target_index - current_index) % option_count
    backward_steps = (current_index - target_index) % option_count
    if forward_steps <= backward_steps:
        return "down", forward_steps
    return "up", backward_steps


def open_schedule_time_popup(
    pyautogui: Any,
    schedule_ui: dict[str, Any],
    config: dict[str, Any],
    args: argparse.Namespace,
    schedule_rect: tuple[int, int, int, int],
    field_name: str,
) -> tuple[int, str, str, tuple[int, int, int, int]]:
    check_macro_cancelled(args)
    process_id = verify_window_target(get_foreground_window_handle(), expected_title="예약 메시지")
    field = schedule_ui.get(field_name, {})
    existing = {handle for handle, _, _, _ in find_kakao_windows()}
    click_window_offset(pyautogui, schedule_rect, float(field.get("x_offset", 288)), float(field.get("y_offset", 450)))
    sleep_delay(config, "after_schedule_field_click", args.delay_scale, minimum=0.15)
    popup = wait_for_kakao_window(exclude_handles=existing, timeout=2.0)
    if popup is None:
        raise FocusSafetyError(f"예약 시각의 {field_name} 선택 창을 찾지 못했습니다.")
    verify_window_target(popup[0], process_id=process_id)
    return popup


def click_schedule_submit(
    pyautogui: Any,
    config: dict[str, Any],
    args: argparse.Namespace,
    schedule_handle: int,
    schedule_rect: tuple[int, int, int, int],
    *,
    before_dispatch: Callable[[], None] | None = None,
) -> None:
    check_macro_cancelled(args)
    if not activate_window_handle(schedule_handle):
        raise FocusSafetyError("예약 메시지 창을 앞으로 가져오지 못했습니다.")
    verify_window_target(schedule_handle, expected_title="예약 메시지")
    # The reviewer may have moved the native form. Never submit with the
    # coordinates captured before that review.
    current_rect = get_window_rect(schedule_handle)
    if current_rect is None:
        raise FocusSafetyError("예약 메시지 창 위치를 다시 확인하지 못해 등록하지 않았습니다.")
    button = config.get("schedule_ui", {}).get("schedule_button", {})
    check_macro_cancelled(args)
    if before_dispatch is not None:
        before_dispatch()
    check_macro_cancelled(args)
    _mark_macro_dispatch("schedule_submit")
    click_window_offset(pyautogui, current_rect, float(button.get("x_offset", 183)), float(button.get("y_offset", 562)))
    sleep_delay(config, "after_schedule_submit", args.delay_scale, minimum=0.8)


LOG_FIELDNAMES = [
    "row_index", "row", "chat_name", "status", "detail", "attachments", "scheduled_at",
    "run_id", "transport", "failure_stage", "retry_transport", "chat_id", "log_id",
]


def write_log(log_dir: Path, results: list[dict[str, Any]], *, run_id: str | None = None) -> Path:
    log_dir.mkdir(parents=True, exist_ok=True)
    unique_id = uuid.uuid4().hex
    path = log_dir / f"send_log_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{unique_id}.csv"
    with path.open("x", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=LOG_FIELDNAMES,
            extrasaction="ignore",
        )
        writer.writeheader()
        writer.writerows({**result, "run_id": run_id or result.get("run_id") or unique_id} for result in results)
    return path


def command_calibrate(args: argparse.Namespace) -> int:
    with pc_kakao_ui_lease(cancel_event=getattr(args, "cancel_event", None), timeout=getattr(args, "pc_lease_timeout", 30.0)):
        check_macro_cancelled(args)
        return _command_calibrate_locked(args)


def _command_calibrate_locked(args: argparse.Namespace) -> int:
    pyautogui = require_pyautogui()
    config = DEFAULT_CONFIG
    if args.config.exists():
        with args.config.open("r", encoding="utf-8") as handle:
            config = deep_merge(DEFAULT_CONFIG, json.load(handle))

    print("Move the mouse pointer to the KakaoTalk conversation search box.")
    input("Keep the pointer there and press Enter in this terminal.")
    position = pyautogui.position()
    config["search_box"] = search_box_config_from_position(int(position.x), int(position.y), list(config.get("app_titles") or []))
    save_config(args.config, config)
    offset = config["search_box"].get("window_offset")
    if offset:
        print(
            f"Saved search box coordinates to {args.config}: "
            f"x={position.x}, y={position.y}, window_offset=({offset['x']}, {offset['y']})"
        )
    else:
        print(f"Saved search box coordinates to {args.config}: x={position.x}, y={position.y}")
    return 0


def command_cursor(args: argparse.Namespace) -> int:
    pyautogui = require_pyautogui()
    print("Showing mouse coordinates. Press Ctrl+C to stop.")
    try:
        while True:
            position = pyautogui.position()
            print(f"x={position.x:4d} y={position.y:4d}", end="\r", flush=True)
            time.sleep(args.interval)
    except KeyboardInterrupt:
        print()
        return 0


def command_preview(args: argparse.Namespace) -> int:
    rows = load_messages(args.file, sheet=args.sheet)
    print(f"Loaded {len(rows)} enabled rows from {args.file}")
    for index, row in enumerate(rows[: args.limit or len(rows)], start=1):
        preview = row.text.replace("\n", "\\n")
        if len(preview) > 80:
            preview = preview[:77] + "..."
        schedule = format_schedule_datetime(row.scheduled_at) or "-"
        print(
            f"{index:03d}. row={row.row_number} chat={row.chat_name} "
            f"attachments={row.attachment_count} scheduled_at={schedule} text={preview}"
        )
    return 0


def command_send(args: argparse.Namespace) -> int:
    with pc_kakao_ui_lease(cancel_event=getattr(args, "cancel_event", None), timeout=getattr(args, "pc_lease_timeout", 30.0)):
        check_macro_cancelled(args)
        return _command_send_locked(args)


def confirm_schedule_ready_cli(row: MessageRow, scheduled_at: datetime) -> bool:
    """Review the actual form while automatic field readback is unavailable."""
    period = "오전" if scheduled_at.hour < 12 else "오후"
    print("\n카카오톡 예약창에 표시된 대상·본문·날짜·오전/오후·시·분을 대조하세요.")
    print(f"대상: {row.chat_name}\n본문:\n{row.text}")
    print(f"예약: {scheduled_at:%Y-%m-%d} {period} {scheduled_at.hour % 12 or 12:02d}:{scheduled_at.minute:02d}")
    print("다르면 카카오톡에서 위 내용으로 수정하세요. 카카오톡의 등록 버튼은 직접 누르지 마세요.")
    try:
        return input("모든 값이 일치하면 확인 입력, 그 외는 중지: ").strip() == "확인"
    except (EOFError, KeyboardInterrupt):
        return False


def _command_send_locked(args: argparse.Namespace) -> int:
    config = load_config(args.config)
    indexed_rows = list(enumerate(load_messages(args.file, sheet=args.sheet)))
    if args.start_at > 1:
        indexed_rows = indexed_rows[args.start_at - 1 :]
    if args.limit:
        indexed_rows = indexed_rows[: args.limit]
    rows = [row for _, row in indexed_rows]

    if not rows:
        print("No rows to send after applying start/limit filters.")
        return 0
    validate_send_mode(len(rows), args.paste_only, args.no_close_chat)

    fallback_schedule = parse_schedule_datetime(args.schedule_at)
    effective_schedules = [schedule_datetime_for_row(row, fallback_schedule) for row in rows]
    scheduled_count = sum(1 for scheduled_at in effective_schedules if scheduled_at is not None)
    if args.paste_only and scheduled_count:
        raise ValueError("예약 메시지는 붙여넣기만 모드와 함께 사용할 수 없습니다.")
    if scheduled_count and not callable(getattr(args, "confirm_schedule_ready", None)):
        args.confirm_schedule_ready = confirm_schedule_ready_cli
    for row, scheduled_at in zip(rows, effective_schedules):
        if scheduled_at is not None:
            validate_schedule_row(row, scheduled_at)

    print(f"Ready to process {len(rows)} row(s) from {args.file}.")
    if scheduled_count:
        print(f"Scheduled message rows: {scheduled_count}")
    print("Use this only for conversations where sending these messages is expected and allowed.")
    if not args.yes:
        confirmation = input("Type SEND to start real KakaoTalk control: ").strip()
        if confirmation != "SEND":
            print("Canceled.")
            return 1

    pyautogui = require_pyautogui()
    pyautogui.PAUSE = args.pause
    countdown(args.countdown, "Starting")

    results: list[dict[str, Any]] = []
    stopped = False
    for index, (row_index, row) in enumerate(indexed_rows, start=1):
        scheduled_at = effective_schedules[index - 1]
        scheduled_at_text = format_schedule_datetime(scheduled_at)
        if args.confirm_each:
            action = f"Schedule {scheduled_at_text}" if scheduled_at else "Send"
            answer = input(f"[{index}/{len(rows)}] {action} to '{row.chat_name}'? Enter=yes, s=skip, q=quit: ").strip().lower()
            if answer == "q":
                results.append(
                    {
                        "row_index": row_index,
                        "row": row.row_number,
                        "chat_name": row.chat_name,
                        "status": "stopped",
                        "detail": "Stopped by user",
                        "attachments": row.attachment_count,
                        "scheduled_at": scheduled_at_text,
                    }
                )
                stopped = True
                break
            if answer == "s":
                results.append(
                    {
                        "row_index": row_index,
                        "row": row.row_number,
                        "chat_name": row.chat_name,
                        "status": "skipped",
                        "detail": "Skipped by user",
                        "attachments": row.attachment_count,
                        "scheduled_at": scheduled_at_text,
                    }
                )
                continue
        try:
            print(f"[{index}/{len(rows)}] Processing row {row.row_number}: {row.chat_name}")
            if scheduled_at is not None:
                registration = send_scheduled_one(pyautogui, config, row, args, scheduled_at)
                verified = type(registration) is ScheduleRegistrationResult and registration.verified is True
                status = "scheduled" if verified else "unknown"
            else:
                send_one(pyautogui, config, row, args)
                status = "sent" if not args.paste_only else "pasted"
            results.append(
                {
                    "row_index": row_index,
                    "row": row.row_number,
                    "chat_name": row.chat_name,
                    "status": status,
                    "detail": (registration.detail if type(registration) is ScheduleRegistrationResult and isinstance(registration.detail, str) else "예약 목록에서 대상·본문·일시 확인 필요") if scheduled_at else "",
                    "attachments": row.attachment_count,
                    "scheduled_at": scheduled_at_text,
                    "transport": "scheduled" if scheduled_at is not None else "macro",
                    "failure_stage": "schedule_verification" if status == "unknown" else "",
                    "retry_transport": "none",
                }
            )
            if status == "unknown":
                print("예약 목록에서 등록 여부를 확인하세요. 중복 등록 방지를 위해 후속 행 처리를 중단합니다.", file=sys.stderr)
                stopped = True
                break
        except KeyboardInterrupt:
            results.append(
                {
                    "row_index": row_index,
                    "row": row.row_number,
                    "chat_name": row.chat_name,
                    "status": "stopped",
                    "detail": "Keyboard interrupt",
                    "attachments": row.attachment_count,
                    "scheduled_at": scheduled_at_text,
                }
            )
            stopped = True
            break
        except Exception as exc:
            results.append(
                {
                    "row_index": row_index,
                    "row": row.row_number,
                    "chat_name": row.chat_name,
                    "status": "unknown" if isinstance(exc, MacroSendUnknownError) else ("stopped" if isinstance(exc, (ScheduleNotSubmittedError, MacroSendCancelled)) else "error"),
                    "detail": str(exc),
                    "failure_stage": getattr(exc, "stage", "schedule_prepare" if scheduled_at is not None else "macro"),
                    "retry_transport": "none",
                    "transport": "scheduled" if scheduled_at is not None else "macro",
                    "attachments": row.attachment_count,
                    "scheduled_at": scheduled_at_text,
                }
            )
            print(f"Error on row {row.row_number}: {exc}", file=sys.stderr)
            if scheduled_at is not None or isinstance(exc, (FocusSafetyError, MacroSendCancelled)) or not args.continue_on_error:
                stopped = True
                break

    if not stopped and not args.paste_only and not args.no_close_chat:
        try:
            return_to_recent_chat_top(pyautogui, config, args)
            print("Returned KakaoTalk recent chat list to the top.")
        except Exception as exc:
            print(f"Warning: Could not return KakaoTalk recent chat list to the top: {exc}", file=sys.stderr)

    run_id = uuid.uuid4().hex
    for result in results:
        result.setdefault("transport", "macro")
    log_path = write_log(args.log_dir, results, run_id=run_id)
    print(f"Wrote log: {log_path}")
    return 0 if not any(row["status"] in {"error", "unknown", "stopped"} for row in results) else 2


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="KakaoTalk CSV/XLSX message sending helper using visible keyboard and mouse control."
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    calibrate = subparsers.add_parser("calibrate", help="Capture the KakaoTalk conversation search box coordinate.")
    calibrate.add_argument("--config", type=Path, default=DEFAULT_CONFIG_PATH)
    calibrate.set_defaults(func=command_calibrate)

    cursor = subparsers.add_parser("cursor", help="Show live mouse coordinates.")
    cursor.add_argument("--interval", type=float, default=0.1)
    cursor.set_defaults(func=command_cursor)

    preview = subparsers.add_parser("preview", help="Preview messages from a CSV or XLSX file.")
    preview.add_argument("--file", type=Path, required=True)
    preview.add_argument("--sheet", default=None, help="XLSX sheet name. Defaults to the active sheet.")
    preview.add_argument("--limit", type=int, default=0)
    preview.set_defaults(func=command_preview)

    send = subparsers.add_parser("send", help="Send or paste messages through KakaoTalk.")
    send.add_argument("--file", type=Path, required=True)
    send.add_argument("--sheet", default=None, help="XLSX sheet name. Defaults to the active sheet.")
    send.add_argument("--config", type=Path, default=DEFAULT_CONFIG_PATH)
    send.add_argument("--yes", action="store_true", help="Skip the initial SEND confirmation.")
    send.add_argument("--confirm-each", action="store_true", help="Ask before each row.")
    send.add_argument("--paste-only", action="store_true", help="Paste the message but do not press Enter to send.")
    send.add_argument("--no-close-chat", action="store_true", help="Do not close the chat window after each row.")
    send.add_argument("--schedule-at", default="", help="Register all selected rows as KakaoTalk scheduled messages. Example: 2026-05-27 09:30")
    send.add_argument("--assume-focused", action="store_true", help="Skip automatic KakaoTalk window activation.")
    send.add_argument("--continue-on-error", action="store_true", help="Continue with the next row after an error.")
    send.add_argument("--start-at", type=int, default=1, help="Start at the Nth loaded message row.")
    send.add_argument("--limit", type=int, default=0, help="Maximum number of rows to process.")
    send.add_argument("--countdown", type=int, default=5, help="Countdown seconds before screen control starts.")
    send.add_argument("--pause", type=float, default=0.05, help="PyAutoGUI pause between low-level actions.")
    send.add_argument("--delay-scale", type=float, default=1.0, help="Multiply configured delays by this value.")
    send.add_argument("--log-dir", type=Path, default=Path("logs"))
    send.set_defaults(func=command_send)

    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return int(args.func(args))
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
